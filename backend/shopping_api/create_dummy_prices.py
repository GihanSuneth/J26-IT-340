
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "data" / "product_recognition" / "product_prices.xlsx"

# Prices and weights are illustrative only.
products = [
    ("anchor_full_cream_milk_powder", "Anchor Full Cream Milk Powder", 1250, 400),
    ("baby_cheramy_moisturising_soap", "Baby Cheramy Moisturising Soap", 280, 100),
    ("coca_cola_pet", "Coca Cola PET", 350, 500),
    ("diva_fresh_rose_and_lime", "Diva Fresh Rose and Lime", 480, 500),
    ("knorr_coconut_milk_powder", "Knorr Coconut Milk Powder", 110, 150),
    ("kotmale_milk_strawberry", "Kotmale Strawberry Milk", 220, 180),
    ("kottumee_noodles", "Kottumee Noodles", 140, 80),
    ("lankasoy_chicken_flavour", "Lankasoy Chicken Flavour", 200, 90),
    ("maggi_coconut_milk", "Maggi Coconut Milk", 300, 150),
    ("maliban_cream_cracker", "Maliban Cream Cracker", 140, 190),
    ("milo_rtd_tetra_pack", "Milo RTD Tetra Pack", 230, 180),
    ("munchee_chocolate_cream", "Munchee Chocolate Cream", 320, 200),
    ("munchee_super_cream_cracker", "Munchee Super Cream Cracker", 350, 190),
    ("my_choice_white_sugar_packet", "My Choice White Sugar", 290, 1000),
    ("pears_soap", "Pears Soap", 260, 100),
    ("raigam_isi_cooking_salt", "Raigam Isi Cooking Salt", 130, 400),
    ("ratthi_full_cream_milk_powder", "Ratthi Full Cream Milk Powder", 1200, 400),
    ("richlife_chocolate_flavoured_milk", "Richlife Chocolate Milk", 230, 180),
    ("sunlight_yellow_detergent_soap", "Sunlight Yellow Detergent Soap", 180, 120),
]

OUTPUT.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Products"

ws.append(["class_name", "product_name", "price_lkr", "expected_weight_g"])

for product in products:
    ws.append(product)

for cell in ws[1]:
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = PatternFill("solid", fgColor="16845A")

for col, width in {"A": 40, "B": 39, "C": 18, "D": 22}.items():
    ws.column_dimensions[col].width = width

ws.freeze_panes = "A2"
wb.save(OUTPUT)

print(f"Created: {OUTPUT}")
print(f"Products: {len(products)}")
