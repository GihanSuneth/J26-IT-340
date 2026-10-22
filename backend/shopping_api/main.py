from pathlib import Path
from threading import Lock, Thread
from collections import deque
import time
import uuid
import cv2
import serial
from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from ultralytics import YOLO

ROOT = Path(__file__).resolve().parents[2]
MODEL = ROOT / 'product model' / 'test' / 'models' / 'best.pt'
PRICES = ROOT / 'data' / 'product_recognition' / 'product_prices.xlsx'
CONFIDENCE = 0.70
CAMERA_INDEX = 1
SERIAL_PORT = 'COM6'
BAUD_RATE = 115200
CAPTURE_SECONDS = 1.0
PRE_EVENT_SECONDS = 0.6
MIN_DETECTION_FRAMES = 2

app = FastAPI(title='SmartCart IR-triggered YOLO API')
app.add_middleware(CORSMiddleware, allow_origins=['http://localhost:5173', 'http://127.0.0.1:5173'], allow_methods=['*'], allow_headers=['*'])
lock = Lock()
products = {}
cart = []
latest_jpeg = None
camera_status = 'starting'
serial_status = 'starting'
last_event = 'none'
pending_events = deque()
detection_history = deque(maxlen=200)


def load_products():
    if not PRICES.exists():
        raise RuntimeError(f'Missing price workbook: {PRICES}')
    wb = load_workbook(PRICES, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    return {str(r[0]): {'class_name': str(r[0]), 'name': str(r[1]), 'price': float(r[2]), 'weight_g': float(r[3])} for r in rows if r[0] is not None}


def cart_snapshot():
    return {'items': [dict(x) for x in cart], 'count': len(cart), 'total': round(sum(x['price'] for x in cart), 2)}


def serial_worker():
    global serial_status
    while True:
        try:
            with serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=0.25) as esp32:
                serial_status = f'connected ({SERIAL_PORT})'
                print('ESP32 connected:', SERIAL_PORT, flush=True)
                while True:
                    line = esp32.readline().decode('utf-8', errors='replace').strip()
                    if line in ('EVENT:ADD', 'EVENT:REMOVE'):
                        now = time.monotonic()
                        with lock:
                            pending_events.append((line.split(':')[1], now))
                        print('IR EVENT:', line, flush=True)
        except (serial.SerialException, OSError) as exc:
            serial_status = f'disconnected: {exc}'
            print('SERIAL:', exc, flush=True)
            time.sleep(2)


def process_events(now):
    global last_event
    # Called while holding lock; allow full 1-second post-event observation.
    while pending_events and now >= pending_events[0][1] + CAPTURE_SECONDS:
        action, event_time = pending_events.popleft()
        start = event_time - PRE_EVENT_SECONDS
        end = event_time + CAPTURE_SECONDS
        frames = [detections for ts, detections in detection_history if start <= ts <= end]
        scores = {}
        for frame_detections in frames:
            for name, confidence in frame_detections.items():
                score = scores.setdefault(name, [0, 0.0])
                score[0] += 1
                score[1] += confidence
        candidates = [(name, count, total / count) for name, (count, total) in scores.items() if count >= MIN_DETECTION_FRAMES]
        candidates.sort(key=lambda x: (x[1], x[2]), reverse=True)
        if not candidates or (len(candidates) > 1 and candidates[0][1] == candidates[1][1]):
            last_event = f'{action}: ignored (no clear product)'
            continue
        name = candidates[0][0]
        if action == 'ADD':
            p = products[name]
            cart.append({'id': str(uuid.uuid4()), 'class_name': name, 'name': p['name'], 'price': p['price'], 'weight_g': p['weight_g'], 'in_basket': True, 'added_at': time.time()})
            last_event = f'ADD: {name}'
        else:
            item = next((item for item in reversed(cart) if item['class_name'] == name), None)
            if item is None:
                last_event = f'REMOVE: ignored ({name} not in cart)'
            else:
                cart.remove(item)
                last_event = f'REMOVE: {name}'
        print('CART EVENT:', last_event, flush=True)


def camera_worker():
    global latest_jpeg, camera_status
    try:
        if not MODEL.is_file():
            raise FileNotFoundError(f'Model not found: {MODEL}')
        model = YOLO(str(MODEL))
        cap = cv2.VideoCapture(CAMERA_INDEX, cv2.CAP_DSHOW)
        if not cap.isOpened():
            cap.release()
            cap = cv2.VideoCapture(CAMERA_INDEX)
        if not cap.isOpened():
            raise RuntimeError('Cannot open webcam. Check CAMERA_INDEX and close other camera apps.')
        cap.set(cv2.CAP_PROP_FRAME_WIDTH, 1280)
        cap.set(cv2.CAP_PROP_FRAME_HEIGHT, 720)
        camera_status = 'running'
        try:
            while True:
                ok, frame = cap.read()
                if not ok:
                    camera_status = 'camera frame unavailable'
                    time.sleep(0.1)
                    continue
                result = model.predict(frame, conf=CONFIDENCE, imgsz=640, iou=0.5, verbose=False)[0]
                detections = {}
                for box in result.boxes:
                    name = str(model.names[int(box.cls.item())])
                    if name in products:
                        detections[name] = max(detections.get(name, 0.0), float(box.conf.item()))
                now = time.monotonic()
                with lock:
                    detection_history.append((now, detections))
                    while detection_history and now - detection_history[0][0] > 6.0:
                        detection_history.popleft()
                    process_events(now)
                annotated = result.plot()
                success, jpeg = cv2.imencode('.jpg', annotated, [int(cv2.IMWRITE_JPEG_QUALITY), 75])
                if success:
                    with lock:
                        latest_jpeg = jpeg.tobytes()
        finally:
            cap.release()
    except Exception as exc:
        camera_status = f'error: {exc}'
        print('CAMERA WORKER ERROR:', exc, flush=True)


@app.on_event('startup')
def startup():
    global products
    products = load_products()
    Thread(target=camera_worker, daemon=True).start()
    Thread(target=serial_worker, daemon=True).start()


@app.get('/api/health')
def health():
    with lock:
        event = last_event
    return {'status': 'ok', 'camera': camera_status, 'serial': serial_status, 'last_event': event, 'model_path': str(MODEL), 'product_count': len(products)}


@app.get('/api/products')
def list_products():
    return list(products.values())


@app.get('/api/cart')
def get_cart():
    with lock:
        return cart_snapshot()


@app.get('/api/video')
def video():
    def stream():
        while True:
            with lock:
                frame = latest_jpeg
            if frame:
                yield b'--frame\r\nContent-Type: image/jpeg\r\n\r\n' + frame + b'\r\n'
            time.sleep(0.05)
    return StreamingResponse(stream(), media_type='multipart/x-mixed-replace; boundary=frame')


@app.delete('/api/cart/{item_id}')
def delete_item(item_id: str):
    with lock:
        item = next((x for x in cart if x['id'] == item_id), None)
        if item is None:
            raise HTTPException(404, 'Item not found')
        if item['in_basket']:
            raise HTTPException(409, 'Physical basket exit must be verified before deletion')
        cart.remove(item)
        return {'status': 'removed', 'cart': cart_snapshot()}


@app.post('/api/cart/reset')
def reset_cart():
    with lock:
        cart.clear()
        pending_events.clear()
        detection_history.clear()
        return cart_snapshot()


if __name__ == '__main__':
    import uvicorn

    uvicorn.run(app, host='127.0.0.1', port=8000)
